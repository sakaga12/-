"""쇼핑몰 주문 엑셀을 택배사 송장 일괄등록 양식으로 바꾼다.

- 같은 사람이 같은 주소로 여러 건 주문하면 한 상자(송장 1장)로 합친다
- 전화번호 형식을 010-0000-0000으로 통일한다
- 취소·반품 주문은 빼고, 주소가 비어 있는 주문은 따로 표시한다

실행: python convert_orders.py  ->  output/송장등록_YYYYMMDD.xlsx
"""
import re
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
SRC = BASE / "input" / "주문내역.xlsx"
OUT = BASE / "output" / f"송장등록_{date.today():%Y%m%d}.xlsx"
EXCLUDE = {"취소", "반품"}


def normalize_phone(raw) -> str:
    digits = re.sub(r"\D", "", str(raw))
    if digits.startswith("82"):
        digits = "0" + digits[2:]
    if len(digits) == 11:
        return f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    return str(raw)


def convert(orders: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    orders = orders[~orders["주문상태"].isin(EXCLUDE)].copy()
    orders["연락처"] = orders["수취인연락처"].map(normalize_phone)
    orders["주소"] = orders["배송지"].fillna("").str.strip()

    problems = orders[orders["주소"] == ""]
    ok = orders[orders["주소"] != ""]
    keys = ["수취인", "연락처", "주소"]

    # 같은 상자에 같은 상품이 여러 번 들어가면 수량을 합친다
    items = (ok.groupby(keys + ["상품명"], sort=False)["수량"].sum().reset_index()
             .assign(품목=lambda d: d["상품명"] + " x" + d["수량"].astype(str))
             .groupby(keys, sort=False)["품목"].agg(" / ".join))

    labels = (ok.groupby(keys, sort=False)
              .agg(총수량=("수량", "sum"),
                   배송메모=("배송메모", lambda s: " / ".join(s.dropna().unique())),
                   주문번호=("주문번호", lambda s: ",".join(map(str, s))))
              .join(items)
              .reset_index()
              .rename(columns={"수취인": "받는분", "연락처": "받는분 전화번호",
                               "주소": "받는분 주소"}))
    labels = labels[["받는분", "받는분 전화번호", "받는분 주소", "품목", "총수량", "배송메모", "주문번호"]]
    return labels, problems[["주문번호", "수취인", "상품명", "주문상태"]]


def save(labels: pd.DataFrame, problems: pd.DataFrame):
    OUT.parent.mkdir(exist_ok=True)
    with pd.ExcelWriter(OUT, engine="openpyxl") as w:
        labels.to_excel(w, sheet_name="송장등록", index=False)
        problems.to_excel(w, sheet_name="확인필요(주소없음)", index=False)

    wb = load_workbook(OUT)
    for ws, color in zip(wb.worksheets, ["1F4E78", "C00000"]):
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=color)
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(width * 1.5, 10), 60)
        ws.freeze_panes = "A2"
    wb.save(OUT)


def main():
    orders = pd.read_excel(SRC)
    labels, problems = convert(orders)
    save(labels, problems)
    print(f"주문 {len(orders)}건 -> 송장 {len(labels)}장 "
          f"(확인 필요 {len(problems)}건) -> {OUT}")


if __name__ == "__main__":
    main()
