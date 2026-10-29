"""books.toscrape.com (크롤링 연습용 공개 사이트) 전체 도서 목록을 엑셀로 수집한다.

실행: python scrape_books.py  ->  output/도서목록.xlsx
요청 사이에 쉬는 시간을 두어 상대 서버에 부담을 주지 않는다.
"""
import time
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

START = "https://books.toscrape.com/catalogue/page-1.html"
DELAY_SEC = 0.5
RATING = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
GBP_TO_KRW = 1850  # 원화 환산 참고용 고정 환율
OUT = Path(__file__).parent / "output" / "도서목록.xlsx"


def scrape() -> list[dict]:
    session = requests.Session()
    session.headers["User-Agent"] = "portfolio-sample-scraper/1.0"
    rows, url = [], START
    while url:
        res = session.get(url, timeout=15)
        res.raise_for_status()
        soup = BeautifulSoup(res.content, "html.parser")
        for book in soup.select("article.product_pod"):
            link = book.h3.a
            price = float(book.select_one(".price_color").text.lstrip("£Â"))
            rows.append({
                "제목": link["title"],
                "가격(GBP)": price,
                "가격(원, 참고)": round(price * GBP_TO_KRW, -1),
                "평점": RATING[book.select_one(".star-rating")["class"][1]],
                "재고": book.select_one(".availability").text.strip(),
                "링크": urljoin(url, link["href"]),
            })
        nxt = soup.select_one("li.next a")
        url = urljoin(url, nxt["href"]) if nxt else None
        print(f"\r{len(rows)}권 수집", end="")
        time.sleep(DELAY_SEC)
    print()
    return rows


def save(rows: list[dict]):
    OUT.parent.mkdir(exist_ok=True)
    df = pd.DataFrame(rows).sort_values(["평점", "가격(GBP)"], ascending=[False, True])
    df.to_excel(OUT, index=False, sheet_name="도서목록")

    wb = load_workbook(OUT)
    ws = wb.active
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E7D32")
    for i, width in enumerate([60, 12, 14, 8, 12, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "£0.00"
        row[2].number_format = "#,##0"
        row[3].number_format = '"★"0'
        row[5].hyperlink = row[5].value
        row[5].font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(OUT)
    print(f"{len(df)}권 저장 -> {OUT}")


if __name__ == "__main__":
    save(scrape())
